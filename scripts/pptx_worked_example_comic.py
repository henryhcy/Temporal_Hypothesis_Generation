"""Editable PowerPoint version of Figure 1, the worked example (2026-10-23, user). It replays the exact drawing code
of scripts/worked_example_figure.py onto a canvas that writes native PowerPoint objects instead of SVG: rectangles,
rounded rectangles, ovals and freeforms for every box, pill, bubble, brace and burst; text boxes for every line;
connectors with arrowheads; the line icons as pictures. Coordinates are printed points, so the slide is the figure's
printed size (240 x 361 pt). The halftone dot pattern of the PDF is omitted.

Fonts: Comic Neue (body) and Bangers (captions), both OFL; copies with their licence are in
docs/draft/Figures/fonts/ -- install them before editing, or PowerPoint substitutes another face.

    .venv313/bin/python scripts/pptx_worked_example_comic.py
"""
from __future__ import annotations

import io
import math
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import worked_example_cartoon as wc  # noqa: E402
import worked_example_figure as wf  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "docs" / "draft" / "Figures" / "worked_example_editable.pptx"
EMU = 12700          # EMU per point


def _rgb(c):
    return RGBColor.from_string(c.lstrip("#").upper())


def _ascent(family, size, weight=400):
    f = ImageFont.truetype(str(wc.FONTDIR / wc.FONTFILES[(family, weight if (family, weight) in wc.FONTFILES else 400)]), 1000)
    a, _ = f.getmetrics()
    return size * a / 1000.0


def _plain(sh):
    sh.shadow.inherit = False
    for ref in sh._element.iter(qn("a:effectRef")):
        ref.set("idx", "0")
    sppr = sh._element.find(qn("p:spPr"))
    if sppr is not None and sppr.find(qn("a:effectLst")) is None:
        etree.SubElement(sppr, qn("a:effectLst"))


class PptxCanvas(wc.Canvas):
    """Same drawing API as worked_example_cartoon.Canvas; writes PowerPoint shapes."""

    def __init__(self, bg):
        super().__init__(bg)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Emu(int(wc.W * EMU)), Emu(int(500 * EMU))
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.sh = self.slide.shapes

    # --- ignored SVG-only bits
    def add(self, s):
        pass

    def _style(self, shape, fill, stroke, sw):
        if fill in (None, "none"):
            shape.fill.background()
        else:
            shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(fill)
        if stroke in (None, "none"):
            shape.line.fill.background()
        else:
            shape.line.color.rgb = _rgb(stroke); shape.line.width = Pt(sw)
        _plain(shape)
        if shape.has_text_frame:
            shape.text_frame.text = ""

    def rect(self, x, y, w, h, r=6, fill="none", stroke="none", sw=1.0, sketch=False, seed=0, dash=None, opacity=1.0):
        if isinstance(fill, str) and fill.startswith("url("):
            return                                      # halftone pattern: PDF only
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE
        s = self.sh.add_shape(kind, Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(w * EMU)), Emu(int(h * EMU)))
        if r:
            s.adjustments[0] = min(0.5, r / min(w, h))
        self._style(s, fill, stroke, sw)

    def circle(self, cx, cy, r, fill, stroke="none", sw=0.8, sketch=False, seed=0):
        s = self.sh.add_shape(MSO_SHAPE.OVAL, Emu(int((cx - r) * EMU)), Emu(int((cy - r) * EMU)), Emu(int(2 * r * EMU)), Emu(int(2 * r * EMU)))
        self._style(s, fill, stroke, sw)

    def path(self, d, fill="none", stroke="#000", sw=1.0, extra=""):
        """Freeform from the absolute M/L/H/V/Q/Z paths the figure code produces (Q sampled)."""
        toks = re.findall(r"[MLHVQZ]|-?\d+(?:\.\d+)?", d.replace(",", " "))
        pts, closed, cur, i = [], False, (0.0, 0.0), 0
        while i < len(toks):
            c = toks[i]; i += 1
            if c in ("M", "L"):
                cur = (float(toks[i]), float(toks[i + 1])); i += 2; pts.append(cur)
            elif c == "H":
                cur = (float(toks[i]), cur[1]); i += 1; pts.append(cur)
            elif c == "V":
                cur = (cur[0], float(toks[i])); i += 1; pts.append(cur)
            elif c == "Q":
                c1 = (float(toks[i]), float(toks[i + 1])); e = (float(toks[i + 2]), float(toks[i + 3])); i += 4
                for t in (0.2, 0.4, 0.6, 0.8, 1.0):
                    pts.append(((1 - t) ** 2 * cur[0] + 2 * (1 - t) * t * c1[0] + t * t * e[0],
                                (1 - t) ** 2 * cur[1] + 2 * (1 - t) * t * c1[1] + t * t * e[1]))
                cur = e
            elif c == "Z":
                closed = True
            else:                                       # implicit L after M
                cur = (float(c), float(toks[i])); i += 1; pts.append(cur)
        if len(pts) < 2:
            return
        b = self.sh.build_freeform(pts[0][0], pts[0][1], scale=EMU)
        b.add_line_segments(pts[1:], close=closed)
        s = b.convert_to_shape()
        self._style(s, fill if closed else "none", stroke, sw)

    def arrow(self, x1, y1, x2, y2, color, sw=1.2, sketch=False, seed=0, head=4.0):
        c = self.sh.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1 * EMU)), Emu(int(y1 * EMU)), Emu(int(x2 * EMU)), Emu(int(y2 * EMU)))
        c.line.color.rgb = _rgb(color); c.line.width = Pt(sw)
        ln = c.line._get_or_add_ln()
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "sm"); t.set("len", "sm")
        _plain(c)

    def _textbox(self, x, y, runs, family, size, anchor="start", italic=False):
        # wrapping box, wide enough never to wrap: LibreOffice centres text in wrap="none" boxes
        width = sum(wc.measure(t, family, size, wt) for t, _, wt in runs) * 1.3 + 10
        top = y - _ascent(family, size)
        left = x - width / 2 if anchor == "middle" else x - 1.0
        tb = self.sh.add_textbox(Emu(int(left * EMU)), Emu(int(top * EMU)), Emu(int(width * EMU)), Emu(int(size * 1.4 * EMU)))
        tf = tb.text_frame; tf.word_wrap = True
        # 0 is dropped as "unset" and PowerPoint falls back to its 0.1 in inset; use 1 EMU
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(1)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if anchor == "middle" else PP_ALIGN.LEFT
        for t, fill, wt in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = family; f.size = Pt(size); f.bold = wt >= 600; f.italic = italic
            f.color.rgb = _rgb(fill)

    def text(self, x, y, s, family, size, fill, weight=400, anchor="start", italic=False, spacing=0):
        self._textbox(x + (1.0 if anchor == "start" else 0), y, [(s, fill, weight)], family, size, anchor, italic)

    def line(self, x, y, runs, family, size, anchor="start"):
        self._textbox(x + 1.0, y, runs, family, size, anchor)

    def picture(self, png, x, y, s):
        self.sh.add_picture(io.BytesIO(png), Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(s * EMU)), Emu(int(s * EMU)))

    def save_pptx(self, H):
        self.prs.slide_height = Emu(int(math.ceil(H) * EMU))
        self.prs.save(OUT)
        print("wrote", OUT, f"({wc.W:.0f} x {H:.0f} pt)")


if __name__ == "__main__":
    cv, H = wf.draw(make=PptxCanvas)
    cv.save_pptx(H)
