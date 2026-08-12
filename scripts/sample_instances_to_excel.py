"""ONE-OFF: render the three sample instances as an .xlsx workbook.

Reads `docs/reference/Sample_Instances.json` (produced by
`scripts/make_sample_instances.py`) and writes
`docs/reference/Sample_Instances.xlsx`.

Five sheets:

- **Overview**   one column per example, fields as rows -- the side-by-side read
- **Evidence**   one row per (example, cohort, set, report_id) -- long/filterable,
                 because the id lists are far too long for single cells
- **Counts**     the per-cell counts behind each argmax / two-cohort test
- **Positional** the descriptive rank statistics, NOT the label
- **MirrorPairs** the supported/refuted pairing per core

Needs `openpyxl` (declared under the `render` extra, not a runtime dependency):

    pip install -e ".[render]"
    python scripts/sample_instances_to_excel.py
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_REPO = Path(__file__).resolve().parent.parent

_HDR = Font(bold=True, color="FFFFFF")
_HDR_FILL = PatternFill("solid", fgColor="44546A")
_SECTION = Font(bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, row=1, ncols=None):
    for c in range(1, (ncols or ws.max_column) + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HDR
        cell.fill = _HDR_FILL
        cell.alignment = Alignment(vertical="top")
    # freeze via the coordinate STRING: ws.cell(...) would CREATE that cell, which
    # bumps max_row and makes every later append() start one row too low, leaving a
    # blank row under each header.
    ws.freeze_panes = f"A{row + 1}"


def _autosize(ws, maxw=70):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), maxw)


def sheet_overview(wb, samples):
    ws = wb.create_sheet("Overview")
    ids = [s["group_id"] for s in samples]
    ws.append(["field"] + ids)
    _style_header(ws, ncols=len(ids) + 1)

    def section(name):
        ws.append([name] + [""] * len(samples))
        for c in range(1, len(samples) + 2):
            ws.cell(row=ws.max_row, column=c).font = _SECTION
            ws.cell(row=ws.max_row, column=c).fill = _SECTION_FILL

    def row(label, getter):
        vals = []
        for s in samples:
            v = getter(s)
            vals.append("" if v is None else
                        (", ".join(map(str, v)) if isinstance(v, list) else v))
        ws.append([label] + vals)
        for c in range(2, len(samples) + 2):
            ws.cell(row=ws.max_row, column=c).alignment = _WRAP

    section("identity")
    for label, key in (("mode", "mode"), ("answer_type", "answer_type"),
                       ("level", "level"), ("why chosen", "why_this_example")):
        row(label, lambda s, k=key: s[k])

    section("hypothesis text")
    row("question (original)", lambda s: s["question_original"])
    for reg in ("original", "research", "media", "clinical"):
        row(f"claim ({reg})", lambda s, r=reg: s["claim_all_registers"][r])

    section("gold label")
    gold_fields = ["gold_verdict", "test_verdict", "insufficient_subtype", "negated",
                   "gold_proportion", "gold_n", "gold_k", "gold_n2", "gold_k2",
                   "ci", "ci_width", "confidence_tag", "effect_size",
                   "options", "gold_answer"]
    for f in gold_fields:
        if any(s["gold"].get(f) is not None for s in samples):
            row(f, lambda s, f=f: s["gold"].get(f))

    section("test")
    for f in ("sig_test", "p_value", "p_claim", "p_opposite", "significant",
              "alpha", "min_n_floor", "family"):
        if any(s["test"].get(f) is not None for s in samples):
            row(f, lambda s, f=f: s["test"].get(f))

    section("evidence sizes (see Evidence sheet for the ids)")

    def sizes(s):
        cc = s["covered_case_ids"]
        if "denominator" in cc:
            return "; ".join(f"{k}={v['count']}" for k, v in cc.items())
        return " | ".join(
            f"{ck}: " + "; ".join(f"{k}={v['count']}" for k, v in cv.items())
            for ck, cv in cc.items())
    row("counts", sizes)
    _autosize(ws)
    ws.column_dimensions["A"].width = 34
    return ws


def sheet_evidence(wb, samples):
    ws = wb.create_sheet("Evidence")
    ws.append(["group_id", "cohort", "set", "report_id"])
    _style_header(ws)
    for s in samples:
        cc = s["covered_case_ids"]
        blocks = ({"-": cc} if "denominator" in cc else cc)
        for cohort, block in blocks.items():
            for setname, payload in block.items():
                for rid in payload["report_ids"]:
                    ws.append([s["group_id"], cohort, setname, rid])
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    _autosize(ws)
    return ws


def sheet_counts(wb, samples):
    ws = wb.create_sheet("Counts")
    ws.append(["group_id", "key", "value"])
    _style_header(ws)
    for s in samples:
        for k, v in (s.get("extra_counts") or {}).items():
            if isinstance(v, dict):
                for kk, vv in v.items():
                    ws.append([s["group_id"], f"{k}.{kk}", vv])
            else:
                ws.append([s["group_id"], k, v])
    _autosize(ws)
    return ws


def sheet_positional(wb, samples):
    ws = wb.create_sheet("Positional")
    ws.append(["group_id", "test", "field", "value"])
    _style_header(ws)
    note = ("DESCRIPTIVE ONLY -- retired as verdict drivers 2026-08-11; the label "
            "comes from the exact test on the Overview sheet")
    ws.append(["(note)", "", "", note])
    ws.cell(row=ws.max_row, column=4).font = Font(italic=True)
    for s in samples:
        for test, payload in (s["positional_statistics_DESCRIPTIVE_ONLY"] or {}).items():
            for k, v in payload.items():
                ws.append([s["group_id"], test, k, v])
    _autosize(ws)
    return ws


def sheet_mirrors(wb, samples):
    ws = wb.create_sheet("MirrorPairs")
    ws.append(["group_id", "derived_from_core_id", "supported text",
               "refuted text", "refuted verdict", "refuted p_claim",
               "refuted p_opposite"])
    _style_header(ws)
    for s in samples:
        mp = s["mirror_pair"]
        ws.append([s["group_id"], mp["derived_from_core_id"], mp["supported_text"],
                   mp["refuted_text"], mp["refuted_verdict"],
                   mp["refuted_p_claim"], mp["refuted_p_opposite"]])
        for c in (3, 4):
            ws.cell(row=ws.max_row, column=c).alignment = _WRAP
    ws.append([])
    ws.append(["Both members of a pair share derived_from_core_id and MUST stay on "
               "the same side of any train/test split."])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True)
    _autosize(ws)
    for c in ("C", "D"):
        ws.column_dimensions[c].width = 60
    return ws


def main() -> None:
    ap = argparse.ArgumentParser(description="Sample instances -> .xlsx")
    ap.add_argument("--src", type=Path,
                    default=_REPO / "docs" / "reference" / "Sample_Instances.json")
    ap.add_argument("--out", type=Path,
                    default=_REPO / "docs" / "reference" / "Sample_Instances.xlsx")
    args = ap.parse_args()

    samples = json.loads(args.src.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    sheet_overview(wb, samples)
    sheet_evidence(wb, samples)
    sheet_counts(wb, samples)
    sheet_positional(wb, samples)
    sheet_mirrors(wb, samples)
    wb.save(args.out)
    print(f"wrote {args.out}")
    for ws in wb.worksheets:
        print(f"  {ws.title:<12} {ws.max_row - 1:>5} data rows x {ws.max_column} cols")


if __name__ == "__main__":
    main()
