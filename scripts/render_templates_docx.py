"""Render the templates document (.docx) from the YAML source files.

Usage:
    python scripts/render_templates_docx.py

Output:
    docs/reference/Temporal_Hypothesis_Templates.docx

The YAML files under data/templates/ are the source of truth. This script
regenerates the human-readable .docx whenever they change.
"""
from pathlib import Path
import yaml
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------- paths ----------------
ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = ROOT / "data" / "templates"
OUTPUT_DIR = ROOT / "docs" / "reference"
OUTPUT_FILE = OUTPUT_DIR / "Temporal_Hypothesis_Templates.docx"

# ---------------- colors ----------------
NAVY  = RGBColor(0x0F, 0x2A, 0x43)
INK   = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)

FAMILY_COLOR = {
    "pairwise":   RGBColor(0x1C, 0x72, 0x93),  # teal
    "sequence":   RGBColor(0x5C, 0x5D, 0x8D),  # indigo
    "first_last": RGBColor(0xC9, 0x7B, 0x4A),  # clay
}
REGISTER_COLOR = {
    "original": RGBColor(0x64, 0x74, 0x8B),  # gray
    "research": RGBColor(0x1C, 0x72, 0x93),  # teal
    "media":    RGBColor(0xC9, 0x7B, 0x4A),  # clay
    "clinical": RGBColor(0x2A, 0x9D, 0x8F),  # green
}
HEADER_FILL = "0F2A43"
ALT_FILL    = "EEF3F6"


# ---------------- helpers ----------------
def shade_cell(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def set_cell_borders(cell, color="CCCCCC", sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), sz)
        b.set(qn("w:color"), color)
        borders.append(b)
    tc_pr.append(borders)


def write_run(paragraph, text, *, font="Arial", size=10, bold=False, italic=False, color=INK):
    r = paragraph.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    return r


def heading(doc, text, level=1, color=NAVY):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        write_run(p, text, size=16, bold=True, color=color)
    else:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        write_run(p, text, size=13, bold=True, color=color)


def body(doc, text, *, italic=False, color=INK, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    write_run(p, text, size=11, italic=italic, color=color)
    return p


# ---------------- group rendering ----------------
def render_group(doc, group, family_name):
    # Group header line
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    write_run(p, group["id"], size=12, bold=True, color=FAMILY_COLOR[family_name])
    write_run(p, "   ·   ", size=11, color=MUTED)
    write_run(p, group["level"].capitalize(), size=11, color=NAVY)
    write_run(p, "   ·   ", size=11, color=MUTED)
    write_run(p, group["answer_type"], size=10, bold=True, color=MUTED)

    # 5-row table (header + 4 register rows)
    table = doc.add_table(rows=5, cols=3)
    table.autofit = False
    table.allow_autofit = False
    widths = [Cm(2.8), Cm(8.0), Cm(8.0)]

    # Header row
    hdr = table.rows[0]
    hdr_cells = hdr.cells
    for i, (txt, w) in enumerate(zip(["Register", "Question", "Claim"], widths)):
        c = hdr_cells[i]
        c.width = w
        shade_cell(c, HEADER_FILL)
        set_cell_borders(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = c.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        write_run(para, txt, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    # Data rows: original, research, media, clinical
    for ri, reg in enumerate(["original", "research", "media", "clinical"]):
        row = table.rows[ri + 1]
        row_cells = row.cells
        fill = ALT_FILL if ri % 2 == 1 else None
        v = group["variants"][reg]

        # Register name cell
        c = row_cells[0]
        c.width = widths[0]
        if fill:
            shade_cell(c, fill)
        set_cell_borders(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = c.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        write_run(para, reg.capitalize(), size=10, bold=True, color=REGISTER_COLOR[reg])

        # Question cell
        c = row_cells[1]
        c.width = widths[1]
        if fill:
            shade_cell(c, fill)
        set_cell_borders(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = c.paragraphs[0]
        write_run(para, v["question"], size=10)

        # Claim cell
        c = row_cells[2]
        c.width = widths[2]
        if fill:
            shade_cell(c, fill)
        set_cell_borders(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = c.paragraphs[0]
        write_run(para, v["claim"], size=10)


# ---------------- main ----------------
def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Title
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    write_run(title, "Temporal Hypothesis Benchmark", size=22, bold=True, color=NAVY)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    write_run(sub, "Question and Claim Templates with Register Projections",
              size=13, color=MUTED)

    body(doc,
         "Generated from data/templates/*.yaml. The YAML files are the source of truth; "
         "this document is regenerated whenever they change.",
         italic=True, color=MUTED, after=14)

    heading(doc, "Registers")
    body(doc, "Each group is projected into three registers plus the canonical "
              "Original seed. All four are generated mechanically via templates and "
              "slot-filling. No LLM rewrites are used.")
    body(doc, "  •  Original: canonical seed, baseline for comparison")
    body(doc, "  •  Research: precise scope, explicit cohort, formal vocabulary")
    body(doc, "  •  Media: public-facing, brand-named, comparison framing")
    body(doc, "  •  Clinical: practitioner-conversational, references to patients", after=14)

    # Per-family sections
    files = [("pairwise.yaml", "3.1  pairwise  family", "pairwise"),
             ("sequence.yaml", "3.2  sequence  family", "sequence"),
             ("first_last.yaml", "3.3  first / last  family", "first_last")]

    heading(doc, "Templates by Family")
    for fname, title_text, key in files:
        path = TEMPLATES_DIR / fname
        with open(path, encoding="utf-8") as f:
            data = yaml.safe_load(f)
        heading(doc, title_text, level=2, color=FAMILY_COLOR[key])
        body(doc, data.get("description", ""), italic=True, color=MUTED, after=4)
        if data.get("notes"):
            body(doc, data["notes"], italic=True, color=MUTED, after=10)
        for grp in data["groups"]:
            render_group(doc, grp, key)

    doc.save(OUTPUT_FILE)
    print(f"Wrote: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
